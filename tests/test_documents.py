"""Documents jobflow writes: the CV as .docx and .tex, the cover letter .docx, and the
interview briefing's HTML."""
import re
from datetime import date

from docx import Document

from jobflow.candidate.documents import write_docx, write_latex, write_letter_docx
from jobflow.candidate.interview import _safe_markup

CV = {
    "name": "José Example", "headline": "Data analyst", "email": "jose@example.com",
    "phone": "tel:+61-400-000-000",
    "sections": {
        "experience": [{"company": "Example Corp", "position": "Data Analyst",
                        "start_date": "2021-02", "end_date": "present", "location": "Sydney",
                        "highlights": [
                            "Built **Python** ETL pipelines processing 2M records daily"]}],
        "education": [{"institution": "Example University", "area": "Statistics", "degree": "BSc",
                       "start_date": "2016", "end_date": "2019"}],
        "publications": [{"title": "Fast Joins", "authors": ["J. Example"], "journal": "VLDB",
                          "date": "2020-05"}],
        "skills": [{"label": "Tools", "details": "Python, SQL"}],
        "awards": ["Dean's list"],
    },
}
EVERY_FACT = ["José Example", "Data analyst", "jose@example.com", "+61-400-000-000",
              "Example Corp", "Data Analyst", "Feb 2021 – present", "Sydney",
              "Built Python ETL pipelines processing 2M records daily",
              "Example University", "BSc in Statistics", "2016 – 2019", "Fast Joins", "J. Example",
              "VLDB", "May 2020", "Tools:", "Python, SQL", "Dean's list"]


def _words(text: str) -> str:
    return re.sub(r"\s+", " ", text)


def test_docx_cv_and_letter_keep_every_word(tmp_path):
    # A .docx missing a fact, or a letter addressed to the wrong person, is sent as it is.
    write_docx(CV, tmp_path / "cv.docx")
    docx_text = _words(" ".join(p.text for p in Document(tmp_path / "cv.docx").paragraphs))
    for fact in EVERY_FACT:
        assert fact in docx_text, fact

    letter = "Dear Jane Smith,\n\nI built ETL pipelines.\n\nKind regards,\nJosé Example"
    write_letter_docx(letter, CV, {"title": "Analyst", "company": "Acme"}, tmp_path / "l.docx",
                      date(2026, 9, 27))
    paragraphs = [p.text for p in Document(tmp_path / "l.docx").paragraphs]
    assert paragraphs[2] == "Jane Smith\nAcme"   # addressed to the greeting's name
    assert paragraphs[-3:] == ["Dear Jane Smith,", "I built ETL pipelines.",
                               "Kind regards,\nJosé Example"]


def test_latex_cv_text_can_never_run_commands(tmp_path):
    # CV text containing LaTeX commands would run them when you compile the .tex, or vanish from the
    # PDF.
    cv = {"name": "A & B", "sections": {"summary": [
        r"**Lead** \input{/etc/passwd} 100% $5 #1 a_b {x} ~ ^", "[Confidential] project"]}}
    write_latex(cv, tmp_path / "cv.tex")
    tex = (tmp_path / "cv.tex").read_text()
    assert r"\textbf{Lead} \textbackslash{}input\{/etc/passwd\} 100\% \$5 \#1 a\_b \{x\} "\
           r"\textasciitilde{} \textasciicircum{}" in tex
    assert r"\input{" not in tex and r"A \& B" in tex
    # A leading [ would otherwise be read as an optional argument and lose its brackets.
    assert "{[}Confidential{]} project" in tex


def test_letter_from_a_non_yaml_cv_takes_its_header_only_from_the_cv(tmp_path):
    # A letter's header from a plain-text CV is guessed: a date or figure taken for a phone number
    # puts a wrong contact detail on it.
    cv_text = ("José Example\njose@example.com | +61 400 000 000\n"
               "Example Corp 2016 - 2019: processed 2,000,000 records")
    letter = "Dear Hiring Manager,\n\nI built ETL pipelines.\n\nKind regards,\nJosé Example"
    write_letter_docx(letter, cv_text, {"title": "Analyst", "company": "Acme"},
                      tmp_path / "l.docx", date(2026, 9, 27))
    paragraphs = [p.text for p in Document(tmp_path / "l.docx").paragraphs]
    # Dates and figures are never taken for a phone number.
    assert paragraphs[0] == "José Example\njose@example.com\n+61 400 000 000"
    write_letter_docx(letter, "José Example\njose@example.com\nAcme (2019 - 2023) analyst",
                      {"title": "Analyst", "company": "Acme"}, tmp_path / "l.docx",
                      date(2026, 9, 27))
    assert Document(tmp_path / "l.docx").paragraphs[0].text == "José Example\njose@example.com"


def test_briefing_markup_cannot_inject_html():
    # The briefing is built from LLM output and web results: script or a javascript: link in it
    # would run when you open the page.
    out = _safe_markup('<script>alert(1)</script>\n'
                       '[click](javascript:alert(1)) [x](https://ok.example/" onmouseover="bad)')
    assert "<script>" not in out
    assert 'href="javascript' not in out and 'onmouseover="bad' not in out
