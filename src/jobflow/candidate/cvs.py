"""Loading your CV files, and helpers shared by matching and tailoring."""
from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path

import yaml

from ..settings import chosen_cvs, cv_directory

JOB_CHAR_LIMIT = 24000
CV_CHAR_LIMIT = 28000
# The most CVs cvs in config/profile.toml may name: scoring makes one LLM call per job
# per CV, each sending the whole CV and ad.
MAX_CVS = 2
CV_FORMATS = (".yaml", ".yml", ".tex", ".docx", ".pdf")
YAML_FORMATS = (".yaml", ".yml")


def variant_name(path: Path) -> str:
    """A CV's short name, which scores are stored under and --cv also takes: the file name
    without its extension and a leading "CV_"."""
    return re.sub(r"^cv[_\- ]+", "", path.stem, flags=re.I) or path.stem


def is_yaml(path: Path) -> bool:
    return path.suffix.casefold() in YAML_FORMATS


def folder_cvs() -> list[str]:
    """The file names in the CV folder in one of CV_FORMATS, for suggesting what cvs in
    config/profile.toml can name. Hidden files and Word's "~$" lock files are left out."""
    directory = cv_directory()
    if not directory.is_dir():
        raise RuntimeError(f"CV folder {directory} not found; create it and put your CVs in "
                           "it, or set paths.cv_directory in config/profile.toml")
    return sorted(p.name for p in directory.iterdir()
                  if p.is_file() and p.suffix.casefold() in CV_FORMATS
                  and not p.name.startswith((".", "~$")))


def chosen_files() -> dict[str, Path]:
    """{short name: file} for the CV files cvs in config/profile.toml names exactly: the
    only CVs scoring, tailor-cv, cover-letter, skill-gaps, and interview-prep use.
    Explains any problem with the setting."""
    names, available = chosen_cvs(), folder_cvs()
    found = ", ".join(available) or f"none found in {cv_directory()}"
    if not names:
        example = ", ".join(f'"{name}"' for name in available[:MAX_CVS]) or '"CV_main.yaml"'
        raise RuntimeError(f"Choose which CVs to use: set cvs = [{example}] at the top of "
                           f"config/profile.toml, naming 1 or {MAX_CVS} CV files exactly "
                           f"(yours: {found})")
    names = list(dict.fromkeys(names))
    if len(names) > MAX_CVS:
        raise RuntimeError(f"cvs in config/profile.toml names {len(names)} CVs "
                           f"({', '.join(names)}); use at most {MAX_CVS}, since each CV is one "
                           "LLM call per job")
    unknown = [name for name in names if name not in available]
    if unknown:
        what = "which isn't a CV file" if len(unknown) == 1 else "which aren't CV files"
        raise RuntimeError(f"cvs in config/profile.toml names {', '.join(unknown)}, {what} "
                           f"in {cv_directory()} (give the exact file name, in one of "
                           f"{', '.join(CV_FORMATS)}); your CV files are: {found}")
    files: dict[str, Path] = {}
    for name in names:
        path = cv_directory() / name
        if variant_name(path) in files:
            raise RuntimeError(f"cvs in config/profile.toml names {files[variant_name(path)].name} "
                               f"and {name}, the same CV in two formats; name one of them")
        files[variant_name(path)] = path
    return files


def chosen_sources() -> dict[str, str]:
    """{short name: CV text} for the CVs cvs in config/profile.toml names."""
    sources = {}
    for variant, path in chosen_files().items():
        stat = path.stat()
        text = cv_text(path, stat.st_mtime_ns, stat.st_size)
        if not text.strip():
            raise RuntimeError(f"No text could be read from {path.name} (a scanned PDF?); "
                               "name another CV in cvs in config/profile.toml")
        sources[variant] = text
    return sources


@lru_cache(maxsize=64)
def cv_text(path: Path, mtime_ns: int = 0, size: int = 0) -> str:
    """A CV file's text: a RenderCV YAML CV's cv section, the body of a .tex without its
    commands, and the text of a .docx or .pdf. Cached until the file changes (mtime_ns and
    size). A file that can't be read (a damaged .docx or .pdf, say) is a RuntimeError naming
    it, not the reading library's own error."""
    try:
        return _read_cv(path)
    except Exception as exc:
        raise RuntimeError(f"{path.name} couldn't be read ({type(exc).__name__}: {exc}); save "
                           "it again, or name another CV in cvs in config/profile.toml") from None


def _read_cv(path: Path) -> str:
    suffix = path.suffix.casefold()
    if suffix in YAML_FORMATS:
        text = path.read_text()
        try:
            document = yaml.safe_load(text)
        except yaml.YAMLError:
            document = None
        # Only the CV itself, as tailor-cv sends it: RenderCV's design, locale, and settings
        # are no evidence, and changing them doesn't make earlier scores out of date.
        if isinstance(document, dict) and isinstance(document.get("cv"), dict):
            text = cv_section_yaml(document)
    elif suffix == ".tex":
        text = path.read_text()
        body = re.search(r"\\begin\{document\}(.*?)\\end\{document\}", text, re.S)
        text = re.sub(r"(?<!\\)%.*", "", body[1] if body else text)
        text = re.sub(r"\\[a-zA-Z]+\*?", " ", text)
        text = re.sub(r"[{}]|\\(?=[^a-zA-Z])", "", text)
    elif suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        lines += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        text = "\n".join(lines)
    else:
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return text


def cv_section_yaml(document: dict) -> str:
    """A RenderCV document's cv section as YAML: what the LLM sees of a YAML CV."""
    return yaml.safe_dump({"cv": document["cv"]}, sort_keys=False, allow_unicode=True)


def job_prompt(job: dict) -> dict:
    """The job as every prompt shows it: title, employer, and the ad, cut to JOB_CHAR_LIMIT."""
    return {"title": job["title"], "company": job["company"],
            "description": job["description"][:JOB_CHAR_LIMIT]}


def check_cv_length(name: str, text: str) -> None:
    """Stop rather than send only the start of a CV, which would silently lose the rest:
    content dropped from a draft, or evidence later in it counted as missing."""
    if len(text) > CV_CHAR_LIMIT:
        raise RuntimeError(f"{name} is {len(text):,} characters; over {CV_CHAR_LIMIT:,} the LLM "
                           "would see only its start, and everything after it would be lost "
                           "or count as missing")
